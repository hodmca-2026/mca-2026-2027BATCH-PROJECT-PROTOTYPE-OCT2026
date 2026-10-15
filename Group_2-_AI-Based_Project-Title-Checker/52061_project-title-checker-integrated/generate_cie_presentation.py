import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
OUTPUT_DIR = r"c:\Users\DELL\project-title-checker\merge\project-title-checker-integrated"
OUTPUT_PPTX = os.path.join(OUTPUT_DIR, "Project_Title_Checker_CIE_II_Final.pptx")
BG_DIR = r"C:\Users\DELL\.gemini\antigravity\brain\9d9cfd09-87de-4116-a0c1-c8aa6293ef61\clean_slides"

# Dimensions: 16:9 Widescreen (13.333 x 7.5 inches)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Colors
COLOR_DARK = RGBColor(0x26, 0x26, 0x26)
COLOR_MUTED = RGBColor(0x55, 0x55, 0x55)
COLOR_GREEN = RGBColor(0x83, 0x99, 0x2A)  # Template divider green
COLOR_ACCENT = RGBColor(0x1B, 0x4D, 0x3E) # Deep academic green
COLOR_RED = RGBColor(0xC0, 0x00, 0x00)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_LIGHT_BG = RGBColor(0xF9, 0xFA, 0xFB)
COLOR_CARD_BORDER = RGBColor(0xDA, 0xE2, 0xC6)
COLOR_BOX_BG1 = RGBColor(0xF4, 0xF7, 0xEE)
COLOR_BOX_BG2 = RGBColor(0xEE, 0xF2, 0xF7)

# Content area boundaries (inside white panel, below thin line):
CONTENT_LEFT = Inches(1.35)
CONTENT_WIDTH = Inches(10.63)

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]

    # Helper: add background image
    def set_slide_bg(slide, bg_num, clean=False):
        if bg_num == 1 and clean:
            bg_path = os.path.join(BG_DIR, "bg_slide1_clean.png")
        else:
            bg_path = os.path.join(BG_DIR, f"bg_slide{bg_num}.png")
        if os.path.exists(bg_path):
            slide.shapes.add_picture(bg_path, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT)

    # Helper: add styled container card
    def add_card(slide, left, top, width, height, bg_color=COLOR_LIGHT_BG, border_color=COLOR_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.2)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: College Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1, 1, clean=True)

    # Clean editable "Prepared by" box
    prep_box = slide1.shapes.add_textbox(Inches(7.2), Inches(6.55), Inches(5.5), Inches(0.75))
    tf1 = prep_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Prepared by : [Student Name / Roll No]"
    p1.font.name = "Times New Roman"
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p1.alignment = PP_ALIGN.RIGHT

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Second Year MCA (2024 Pattern) — Semester III"
    p1_sub.font.name = "Times New Roman"
    p1_sub.font.size = Pt(12)
    p1_sub.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    p1_sub.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # SLIDE 2: TITLE OF PROTOTYPE
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2, 2)

    # Prototype Main Title Box
    title_box = slide2.shapes.add_textbox(CONTENT_LEFT, Inches(2.78), CONTENT_WIDTH, Inches(1.15))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "AI-Based Project Title and Proposal Similarity Detection System"
    p2.font.name = "Times New Roman"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT
    p2.alignment = PP_ALIGN.CENTER

    p2_sub = tf2.add_paragraph()
    p2_sub.text = "An Academic Intelligence Platform Integrating NLP, Deep Learning & Information Retrieval"
    p2_sub.font.name = "Times New Roman"
    p2_sub.font.size = Pt(15)
    p2_sub.font.italic = True
    p2_sub.font.color.rgb = COLOR_MUTED
    p2_sub.alignment = PP_ALIGN.CENTER

    # 3 Summary Pillar Cards
    card_w = Inches(3.38)
    card_gap = Inches(0.24)
    card_top = Inches(4.05)
    card_h = Inches(2.70)

    pillars = [
        ("Academic Motivation & Scope", COLOR_BOX_BG1, [
            "Problem: Students frequently propose duplicate or cosmetically rephrased project titles.",
            "Challenge: Manual verification across past cohorts is time-consuming and prone to human oversight.",
            "Solution: Multi-stage verification analyzing both proposal Title and detailed Methodology."
        ]),
        ("Core Technical Approach", COLOR_BOX_BG2, [
            "Hybrid Similarity Engine: Combines MiniLM-L6-v2 semantic embeddings with TF-IDF cosine & Jaccard metrics.",
            "Deep Learning Verification: BiLSTM sequence domain classification, Autoencoder novelty detection & Neural NER.",
            "Dual-Threshold Governance: Automated duplicate prevention (>60%) with faculty review escalation."
        ]),
        ("Role in Decision Support", COLOR_BOX_BG1, [
            "Faculty Autonomy: AI provides calibrated advisory signals and similarity evidence, not automated verdicts.",
            "Uncertainty Handling: Flags ambiguous or out-of-distribution topics as 'Needs Faculty Review'.",
            "Actionable Governance: Interactive review modal with one-click approval or corrective student warning alerts."
        ])
    ]

    for idx, (col_title, bg_col, points) in enumerate(pillars):
        c_left = CONTENT_LEFT + idx * (card_w + card_gap)
        add_card(slide2, c_left, card_top, card_w, card_h, bg_col)
        tb = slide2.shapes.add_textbox(c_left + Inches(0.15), card_top + Inches(0.12), card_w - Inches(0.3), card_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = col_title
        pt.font.name = "Times New Roman"
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_ACCENT
        pt.space_after = Pt(6)

        for pt_text in points:
            p = tf.add_paragraph()
            p.text = pt_text
            p.font.name = "Times New Roman"
            p.font.size = Pt(11.5)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 3: Architecture Diagram (Flow of system)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3, 3)

    sub3 = slide3.shapes.add_textbox(CONTENT_LEFT, Inches(2.72), CONTENT_WIDTH, Inches(0.35))
    tf3_sub = sub3.text_frame
    tf3_sub.word_wrap = True
    p3_sub = tf3_sub.paragraphs[0]
    p3_sub.text = "End-to-End System Architecture: Verification, Deep Learning Analysis & Decision Flow"
    p3_sub.font.name = "Times New Roman"
    p3_sub.font.size = Pt(13)
    p3_sub.font.italic = True
    p3_sub.font.color.rgb = COLOR_MUTED
    p3_sub.alignment = PP_ALIGN.CENTER

    flow_top = Inches(3.12)
    col1_w = Inches(2.2)
    add_card(slide3, CONTENT_LEFT, flow_top, col1_w, Inches(3.6), COLOR_BOX_BG2)
    tb_c1 = slide3.shapes.add_textbox(CONTENT_LEFT + Inches(0.1), flow_top + Inches(0.12), col1_w - Inches(0.2), Inches(3.36))
    tfc1 = tb_c1.text_frame
    tfc1.word_wrap = True
    pc1 = tfc1.paragraphs[0]
    pc1.text = "1. Student Submission"
    pc1.font.name = "Times New Roman"
    pc1.font.size = Pt(13)
    pc1.font.bold = True
    pc1.font.color.rgb = COLOR_ACCENT
    pc1.space_after = Pt(6)
    for t in [
        "Web Interface: Clean student submission form (/student).",
        "Dual Inputs: Title + Technical Description / Abstract.",
        "Role Auth: Student login with session verification.",
        "DB Storage: SQLite schema stores title, status, timestamp."
    ]:
        p = tfc1.add_paragraph()
        p.text = t
        p.font.name = "Times New Roman"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # Arrow 1 -> 2
    arr1 = slide3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, CONTENT_LEFT + col1_w + Inches(0.05), flow_top + Inches(1.6), Inches(0.28), Inches(0.3))
    arr1.fill.solid(); arr1.fill.fore_color.rgb = COLOR_GREEN; arr1.line.fill.background()

    # Flow Box 2: Preprocessing & Routing
    col2_left = CONTENT_LEFT + col1_w + Inches(0.38)
    col2_w = Inches(2.3)
    add_card(slide3, col2_left, flow_top, col2_w, Inches(3.6), COLOR_BOX_BG1)
    tb_c2 = slide3.shapes.add_textbox(col2_left + Inches(0.1), flow_top + Inches(0.12), col2_w - Inches(0.2), Inches(3.36))
    tfc2 = tb_c2.text_frame
    tfc2.word_wrap = True
    pc2 = tfc2.paragraphs[0]
    pc2.text = "2. Flask API & Prep"
    pc2.font.name = "Times New Roman"
    pc2.font.size = Pt(13)
    pc2.font.bold = True
    pc2.font.color.rgb = COLOR_ACCENT
    pc2.space_after = Pt(6)
    for t in [
        "RESTful Endpoints: Flask backend routes in app.py.",
        "Text Normalization: Lowercasing, punctuation stripping.",
        "Tokenization & Stemming: NLTK Porter Stemmer + Vocab mapping.",
        "Payload Routing: Parallel dispatch to IR and DL engines."
    ]:
        p = tfc2.add_paragraph()
        p.text = t
        p.font.name = "Times New Roman"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # Arrow 2 -> 3
    arr2 = slide3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, col2_left + col2_w + Inches(0.05), flow_top + Inches(1.6), Inches(0.28), Inches(0.3))
    arr2.fill.solid(); arr2.fill.fore_color.rgb = COLOR_GREEN; arr2.line.fill.background()

    # Flow Box 3: Dual Analysis Engines (IR + DL)
    col3_left = col2_left + col2_w + Inches(0.38)
    col3_w = Inches(3.0)
    add_card(slide3, col3_left, flow_top, col3_w, Inches(3.6), COLOR_WHITE, COLOR_ACCENT)
    tb_c3 = slide3.shapes.add_textbox(col3_left + Inches(0.1), flow_top + Inches(0.12), col3_w - Inches(0.2), Inches(3.36))
    tfc3 = tb_c3.text_frame
    tfc3.word_wrap = True
    pc3 = tfc3.paragraphs[0]
    pc3.text = "3. Core AI Analysis Engines"
    pc3.font.name = "Times New Roman"
    pc3.font.size = Pt(13)
    pc3.font.bold = True
    pc3.font.color.rgb = COLOR_ACCENT
    pc3.space_after = Pt(4)

    p_ir = tfc3.add_paragraph()
    p_ir.text = "[A] Similarity & IR Engine (ml_engine.py):"
    p_ir.font.name = "Times New Roman"
    p_ir.font.size = Pt(11)
    p_ir.font.bold = True
    p_ir.font.color.rgb = COLOR_DARK
    for t in [
        "• Semantic: all-MiniLM-L6-v2 transformer cosine similarity.",
        "• Lexical: TF-IDF n-grams (0.45) + Jaccard (0.30) + Sequence (0.25)."
    ]:
        p = tfc3.add_paragraph()
        p.text = t
        p.font.name = "Times New Roman"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK

    p_dl = tfc3.add_paragraph()
    p_dl.text = "[B] Deep Learning Engine (dl_engine/):"
    p_dl.font.name = "Times New Roman"
    p_dl.font.size = Pt(11)
    p_dl.font.bold = True
    p_dl.font.color.rgb = COLOR_DARK
    p_dl.space_before = Pt(4)
    for t in [
        "• Unit IV: BiLSTM Domain Classifier + Calibrated LogReg.",
        "• Unit V: Dense Autoencoder Reconstruction & Novelty MSE.",
        "• Unit VI: BiLSTM Neural NER + Technology Gazetteer."
    ]:
        p = tfc3.add_paragraph()
        p.text = t
        p.font.name = "Times New Roman"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK

    # Arrow 3 -> 4
    arr3 = slide3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, col3_left + col3_w + Inches(0.05), flow_top + Inches(1.6), Inches(0.28), Inches(0.3))
    arr3.fill.solid(); arr3.fill.fore_color.rgb = COLOR_GREEN; arr3.line.fill.background()

    # Flow Box 4: Faculty Decision Dashboard
    col4_left = col3_left + col3_w + Inches(0.38)
    col4_w = Inches(2.29)
    add_card(slide3, col4_left, flow_top, col4_w, Inches(3.6), COLOR_BOX_BG2)
    tb_c4 = slide3.shapes.add_textbox(col4_left + Inches(0.1), flow_top + Inches(0.12), col4_w - Inches(0.2), Inches(3.36))
    tfc4 = tb_c4.text_frame
    tfc4.word_wrap = True
    pc4 = tfc4.paragraphs[0]
    pc4.text = "4. Faculty Review & Governance"
    pc4.font.name = "Times New Roman"
    pc4.font.size = Pt(13)
    pc4.font.bold = True
    pc4.font.color.rgb = COLOR_ACCENT
    pc4.space_after = Pt(6)
    for t in [
        "Risk Stratification: Low Risk (<45%), Review (45-60%), High Risk (>=60%).",
        "DL Review Modal: Domain uncertainty, Novelty MSE, extracted tech stack.",
        "Governance Actions: One-click Approve or issue Warning Feedback with guidance.",
        "Audit Trail: Updated status and alerts pushed to student dashboard."
    ]:
        p = tfc4.add_paragraph()
        p.text = t
        p.font.name = "Times New Roman"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 4: Contents of Unit IV
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4, 4)

    sub4 = slide4.shapes.add_textbox(CONTENT_LEFT, Inches(2.72), CONTENT_WIDTH, Inches(0.35))
    tf4_sub = sub4.text_frame
    tf4_sub.word_wrap = True
    p4_sub = tf4_sub.paragraphs[0]
    p4_sub.text = "Unit IV: Recurrent Neural Networks (RNN) & LSTM for Sequence Learning"
    p4_sub.font.name = "Times New Roman"
    p4_sub.font.size = Pt(14)
    p4_sub.font.italic = True
    p4_sub.font.color.rgb = COLOR_ACCENT
    p4_sub.alignment = PP_ALIGN.CENTER

    col_w2 = Inches(5.15)
    gap2 = Inches(0.33)
    top2 = Inches(3.15)
    h2 = Inches(3.62)

    # Left Card: Syllabus Theory
    add_card(slide4, CONTENT_LEFT, top2, col_w2, h2, COLOR_BOX_BG1)
    tb_u4_left = slide4.shapes.add_textbox(CONTENT_LEFT + Inches(0.2), top2 + Inches(0.12), col_w2 - Inches(0.4), h2 - Inches(0.24))
    tf_u4_l = tb_u4_left.text_frame
    tf_u4_l.word_wrap = True
    pu4_l_title = tf_u4_l.paragraphs[0]
    pu4_l_title.text = "Unit IV Syllabus Theoretical Framework"
    pu4_l_title.font.name = "Times New Roman"
    pu4_l_title.font.size = Pt(15.5)
    pu4_l_title.font.bold = True
    pu4_l_title.font.color.rgb = COLOR_ACCENT
    pu4_l_title.space_after = Pt(6)

    u4_theory_points = [
        ("Sequence Learning Problem:", "Standard feedforward networks fail to capture temporal word dependencies in variable-length proposals."),
        ("Recurrent Neural Networks (RNNs):", "Process sequential tokens by passing hidden state vectors through unfolded computational graphs over time."),
        ("Vanishing / Exploding Gradients:", "Mitigated through gating mechanisms that regulate error gradient flow across long recurrent chains."),
        ("Long Short-Term Memory (LSTM):", "Maintains cell state via selective read, selective write, and selective forget gates to preserve long-range dependencies."),
        ("Bidirectional RNNs (BiLSTM):", "Processes sequences forward and backward, capturing preceding and succeeding contextual semantics.")
    ]
    for heading, desc in u4_theory_points:
        p = tf_u4_l.add_paragraph()
        run_h = p.add_run()
        run_h.text = f"• {heading} "
        run_h.font.bold = True
        run_h.font.size = Pt(10.5)
        run_h.font.color.rgb = COLOR_DARK
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(10.5)
        run_d.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

    # Right Card: Prototype Implementation
    u4_r_left = CONTENT_LEFT + col_w2 + gap2
    add_card(slide4, u4_r_left, top2, col_w2, h2, COLOR_BOX_BG2)
    tb_u4_right = slide4.shapes.add_textbox(u4_r_left + Inches(0.2), top2 + Inches(0.12), col_w2 - Inches(0.4), h2 - Inches(0.24))
    tf_u4_r = tb_u4_right.text_frame
    tf_u4_r.word_wrap = True
    pu4_r_title = tf_u4_r.paragraphs[0]
    pu4_r_title.text = "Prototype Implementation & Engineering (dl_engine)"
    pu4_r_title.font.name = "Times New Roman"
    pu4_r_title.font.size = Pt(15.5)
    pu4_r_title.font.bold = True
    pu4_r_title.font.color.rgb = COLOR_ACCENT
    pu4_r_title.space_after = Pt(6)

    u4_impl_points = [
        ("Architecture (ProjectDomainLSTM):", "PyTorch Bidirectional LSTM with 64-dim learned embeddings, hidden_dim=64, 2 layers, and dropout=0.3."),
        ("Sequence Modeling Pipeline:", "Vocabulary builds token mapping (<pad>=0, <unk>=1); encodes variable-length proposal text padded/truncated to max_len=60 tokens."),
        ("7 Domain Classes:", "Predicts Healthcare, Education, Finance, Agriculture, Cybersecurity, Web Applications, and Other domains."),
        ("Calibrated Dual-Model Strategy:", "TF-IDF + Calibrated Logistic Regression serves as primary classifier, backed by the PyTorch BiLSTM sequence model."),
        ("Empirical Evaluation & Safety:", "Both models evaluated on an identical test split. Conservative thresholding flags low-confidence topics as 'Needs Faculty Review'.")
    ]
    for heading, desc in u4_impl_points:
        p = tf_u4_r.add_paragraph()
        run_h = p.add_run()
        run_h.text = f"• {heading} "
        run_h.font.bold = True
        run_h.font.size = Pt(10.5)
        run_h.font.color.rgb = COLOR_DARK
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(10.5)
        run_d.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 5: Contents of Unit V
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5, 5)

    sub5 = slide5.shapes.add_textbox(CONTENT_LEFT, Inches(2.72), CONTENT_WIDTH, Inches(0.35))
    tf5_sub = sub5.text_frame
    tf5_sub.word_wrap = True
    p5_sub = tf5_sub.paragraphs[0]
    p5_sub.text = "Unit V: Autoencoders & Deep Representations for Anomaly and Novelty Detection"
    p5_sub.font.name = "Times New Roman"
    p5_sub.font.size = Pt(14)
    p5_sub.font.italic = True
    p5_sub.font.color.rgb = COLOR_ACCENT
    p5_sub.alignment = PP_ALIGN.CENTER

    # Left Card: Syllabus Theory
    add_card(slide5, CONTENT_LEFT, top2, col_w2, h2, COLOR_BOX_BG1)
    tb_u5_left = slide5.shapes.add_textbox(CONTENT_LEFT + Inches(0.2), top2 + Inches(0.12), col_w2 - Inches(0.4), h2 - Inches(0.24))
    tf_u5_l = tb_u5_left.text_frame
    tf_u5_l.word_wrap = True
    pu5_l_title = tf_u5_l.paragraphs[0]
    pu5_l_title.text = "Unit V Syllabus Theoretical Framework"
    pu5_l_title.font.name = "Times New Roman"
    pu5_l_title.font.size = Pt(15.5)
    pu5_l_title.font.bold = True
    pu5_l_title.font.color.rgb = COLOR_ACCENT
    pu5_l_title.space_after = Pt(6)

    u5_theory_points = [
        ("Architecture of Autoencoders:", "Unsupervised networks with an Encoder compressing input x into bottleneck z, and a Decoder reconstructing x' from z."),
        ("Bottleneck Representation:", "Compresses input through lower-dimensional latent layers (dim(z) << dim(x)), forcing learning of salient regularities."),
        ("Types of Autoencoders:", "Undercomplete, Denoising (noise injection), Variational (probabilistic priors), and Deep/Stacked Autoencoders."),
        ("Loss Function (Reconstruction MSE):", "Optimized using Mean Squared Error: L(x, x') = ||x - x'||^2 between original and reconstructed feature vectors."),
        ("Syllabus Case Study Alignment:", "Applied in anomaly and novelty estimation: typical inputs reconstruct with low MSE, while anomalous proposals yield high MSE.")
    ]
    for heading, desc in u5_theory_points:
        p = tf_u5_l.add_paragraph()
        run_h = p.add_run()
        run_h.text = f"• {heading} "
        run_h.font.bold = True
        run_h.font.size = Pt(10.5)
        run_h.font.color.rgb = COLOR_DARK
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(10.5)
        run_d.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

    # Right Card: Prototype Implementation
    u5_r_left = CONTENT_LEFT + col_w2 + gap2
    add_card(slide5, u5_r_left, top2, col_w2, h2, COLOR_BOX_BG2)
    tb_u5_right = slide5.shapes.add_textbox(u5_r_left + Inches(0.2), top2 + Inches(0.12), col_w2 - Inches(0.4), h2 - Inches(0.24))
    tf_u5_r = tb_u5_right.text_frame
    tf_u5_r.word_wrap = True
    pu5_r_title = tf_u5_r.paragraphs[0]
    pu5_r_title.text = "Prototype Implementation & Engineering (dl_engine)"
    pu5_r_title.font.name = "Times New Roman"
    pu5_r_title.font.size = Pt(15.5)
    pu5_r_title.font.bold = True
    pu5_r_title.font.color.rgb = COLOR_ACCENT
    pu5_r_title.space_after = Pt(6)

    u5_impl_points = [
        ("PyTorch Model (ProjectAutoencoder):", "Dense architecture: Input (100 TF-IDF features) → 48-dim Encoder → 24-dim Bottleneck → 48-dim Decoder → 100-dim Output."),
        ("Training & Artifacts:", "Trained on academic project proposal distributions; weights saved in best_autoencoder_model.pt with feature vocabulary autoencoder_vectorizer.pkl."),
        ("Reconstruction MSE Scoring:", "Computes test reconstruction error calibrated against training distribution (baseline MSE ~0.0075, std ~0.0076)."),
        ("Calibrated Novelty Signals:", "Low MSE indicates typical project patterns; reconstruction error > threshold flags unusual/novel topics for faculty review."),
        ("Zero-Vector & OOD Safeguards:", "If input text shares zero vocabulary with training corpus, gracefully flags 'Unfamiliar Outlier' rather than reporting false typicality.")
    ]
    for heading, desc in u5_impl_points:
        p = tf_u5_r.add_paragraph()
        run_h = p.add_run()
        run_h.text = f"• {heading} "
        run_h.font.bold = True
        run_h.font.size = Pt(10.5)
        run_h.font.color.rgb = COLOR_DARK
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(10.5)
        run_d.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 6: Contents of Unit VI
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6, 6)

    sub6 = slide6.shapes.add_textbox(CONTENT_LEFT, Inches(2.72), CONTENT_WIDTH, Inches(0.35))
    tf6_sub = sub6.text_frame
    tf6_sub.word_wrap = True
    p6_sub = tf6_sub.paragraphs[0]
    p6_sub.text = "Unit VI: Applications of Deep Learning — Named Entity Recognition & Project Analytics"
    p6_sub.font.name = "Times New Roman"
    p6_sub.font.size = Pt(14)
    p6_sub.font.italic = True
    p6_sub.font.color.rgb = COLOR_ACCENT
    p6_sub.alignment = PP_ALIGN.CENTER

    # Left Card: Syllabus Applications
    add_card(slide6, CONTENT_LEFT, top2, col_w2, h2, COLOR_BOX_BG1)
    tb_u6_left = slide6.shapes.add_textbox(CONTENT_LEFT + Inches(0.2), top2 + Inches(0.12), col_w2 - Inches(0.4), h2 - Inches(0.24))
    tf_u6_l = tb_u6_left.text_frame
    tf_u6_l.word_wrap = True
    pu6_l_title = tf_u6_l.paragraphs[0]
    pu6_l_title.text = "Unit VI Syllabus Theoretical Applications"
    pu6_l_title.font.name = "Times New Roman"
    pu6_l_title.font.size = Pt(15.5)
    pu6_l_title.font.bold = True
    pu6_l_title.font.color.rgb = COLOR_ACCENT
    pu6_l_title.space_after = Pt(6)

    u6_theory_points = [
        ("Project Pipeline Design:", "Structuring complete real-world deep learning pipelines: data ingestion, preprocessing, inference, evaluation, and user delivery."),
        ("Named Entity Recognition (NER):", "Sequential token classification to locate and classify key technical units into predefined categories from unstructured text."),
        ("Recommender Systems in DL:", "Content-based and latent vector recommendation models to identify related items based on semantic proximity."),
        ("Multi-Modal & NLP Case Studies:", "Syllabus covers object detection, facial recognition, image captioning, self-driving vehicles, and automated entity extraction."),
        ("Integration into Operational Systems:", "Translating neural inference outputs into interactive dashboards with explainable indicators for domain specialists.")
    ]
    for heading, desc in u6_theory_points:
        p = tf_u6_l.add_paragraph()
        run_h = p.add_run()
        run_h.text = f"• {heading} "
        run_h.font.bold = True
        run_h.font.size = Pt(10.5)
        run_h.font.color.rgb = COLOR_DARK
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(10.5)
        run_d.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

    # Right Card: Prototype Implementation
    u6_r_left = CONTENT_LEFT + col_w2 + gap2
    add_card(slide6, u6_r_left, top2, col_w2, h2, COLOR_BOX_BG2)
    tb_u6_right = slide6.shapes.add_textbox(u6_r_left + Inches(0.2), top2 + Inches(0.12), col_w2 - Inches(0.4), h2 - Inches(0.24))
    tf_u6_r = tb_u6_right.text_frame
    tf_u6_r.word_wrap = True
    pu6_r_title = tf_u6_r.paragraphs[0]
    pu6_r_title.text = "Prototype Implementation & Engineering (dl_engine)"
    pu6_r_title.font.name = "Times New Roman"
    pu6_r_title.font.size = Pt(15.5)
    pu6_r_title.font.bold = True
    pu6_r_title.font.color.rgb = COLOR_ACCENT
    pu6_r_title.space_after = Pt(6)

    u6_impl_points = [
        ("Neural Token Classifier (NeuralNERModel):", "BiLSTM sequence tagger with 48-dim embeddings, hidden_dim=48, and linear classifier predicting B-TECH, I-TECH, and O tags."),
        ("Technical Domain Gazetteer:", "High-precision rule-based gazetteer categorizing extracted entities into Languages (Python, Java), Frameworks (PyTorch, Flask), and Methods (CNN, LSTM)."),
        ("Noise Filtering & Post-Processing:", "Regex guards strip out generic stopwords, numeric fragments, and common adjectives to preserve clean technical entities."),
        ("Content-Based Project Recommendation:", "TF-IDF + Cosine similarity recommender surfaces top-3 semantically related approved proposals for any new idea."),
        ("Faculty DL Review Modal UI:", "Presents extracted tech stack, novelty indicators, similarity matches, and calibrated domain badges in a unified modal.")
    ]
    for heading, desc in u6_impl_points:
        p = tf_u6_r.add_paragraph()
        run_h = p.add_run()
        run_h.text = f"• {heading} "
        run_h.font.bold = True
        run_h.font.size = Pt(10.5)
        run_h.font.color.rgb = COLOR_DARK
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(10.5)
        run_d.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 7: Conclusion
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7, 7)

    sub7 = slide7.shapes.add_textbox(CONTENT_LEFT, Inches(2.72), CONTENT_WIDTH, Inches(0.35))
    tf7_sub = sub7.text_frame
    tf7_sub.word_wrap = True
    p7_sub = tf7_sub.paragraphs[0]
    p7_sub.text = "System Synthesis: Academic Impact, Operational Boundaries & Decision Support"
    p7_sub.font.name = "Times New Roman"
    p7_sub.font.size = Pt(14)
    p7_sub.font.italic = True
    p7_sub.font.color.rgb = COLOR_ACCENT
    p7_sub.alignment = PP_ALIGN.CENTER

    grid_w = Inches(5.15)
    grid_h = Inches(1.72)
    grid_gap_x = Inches(0.33)
    grid_gap_y = Inches(0.18)
    top_row_y = Inches(3.15)
    bot_row_y = top_row_y + grid_h + grid_gap_y

    concl_cards = [
        ("Core Problem Addressed", CONTENT_LEFT, top_row_y, COLOR_BOX_BG1, [
            "Effectively mitigates redundant, duplicated, or cosmetically rephrased student project titles.",
            "Moves beyond simple keyword overlap by analyzing full proposal context (Title + Description)."
        ]),
        ("Synergy of IR & Deep Learning", CONTENT_LEFT + grid_w + grid_gap_x, top_row_y, COLOR_BOX_BG2, [
            "Information Retrieval provides fast, explainable token ranking and content recommendation.",
            "Deep Learning adds semantic sentence embeddings, LSTM domain profiling, and Autoencoder novelty checks."
        ]),
        ("Value for Faculty Decision-Making", CONTENT_LEFT, bot_row_y, COLOR_BOX_BG2, [
            "Replaces manual keyword searching with an automated similarity risk assessment.",
            "Delivers an evidence-backed DL review modal: extracted tech stack, closest matches, and domain confidence."
        ]),
        ("System Limitations & Human Governance", CONTENT_LEFT + grid_w + grid_gap_x, bot_row_y, COLOR_BOX_BG1, [
            "AI serves strictly as decision support; high similarity indicates overlap, not proof of plagiarism.",
            "Faculty retain final authority to approve, reject, or request methodology modifications."
        ])
    ]

    for title, left, top, bg_col, pts in concl_cards:
        add_card(slide7, left, top, grid_w, grid_h, bg_col)
        tb = slide7.shapes.add_textbox(left + Inches(0.15), top + Inches(0.10), grid_w - Inches(0.3), grid_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Times New Roman"
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_ACCENT
        pt.space_after = Pt(4)
        for t in pts:
            p = tf.add_paragraph()
            p.text = f"• {t}"
            p.font.name = "Times New Roman"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 8: Future Enhancement
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide8, 8)

    sub8 = slide8.shapes.add_textbox(CONTENT_LEFT, Inches(2.72), CONTENT_WIDTH, Inches(0.35))
    tf8_sub = sub8.text_frame
    tf8_sub.word_wrap = True
    p8_sub = tf8_sub.paragraphs[0]
    p8_sub.text = "Roadmap for Continued Research, Model Refinement & Scalability"
    p8_sub.font.name = "Times New Roman"
    p8_sub.font.size = Pt(14)
    p8_sub.font.italic = True
    p8_sub.font.color.rgb = COLOR_ACCENT
    p8_sub.alignment = PP_ALIGN.CENTER

    future_cards = [
        ("1. Dataset Expansion & Benchmarking", CONTENT_LEFT, top_row_y, COLOR_BOX_BG1, [
            "Expand beyond illustrative academic examples to multi-year, multi-institutional thesis repositories.",
            "Establish domain-specific annotated test benchmarks for MCA project evaluation."
        ]),
        ("2. Fine-Tuned Transformer Embeddings", CONTENT_LEFT + grid_w + grid_gap_x, top_row_y, COLOR_BOX_BG2, [
            "Explore domain fine-tuning of sentence transformers (e.g. SciBERT) on technical computer science abstracts.",
            "Improve semantic cross-encoder re-ranking for complex paraphrase detection."
        ]),
        ("3. Advanced Neural NER & Taxonomy", CONTENT_LEFT, bot_row_y, COLOR_BOX_BG2, [
            "Transition from BiLSTM token tagging to Transformer-based token classification (BERT-NER).",
            "Incorporate broader ontology mappings for evolving AI frameworks and cloud architectures."
        ]),
        ("4. Institutional Multi-Department Governance", CONTENT_LEFT + grid_w + grid_gap_x, bot_row_y, COLOR_BOX_BG1, [
            "Support department-level custom domains, guide assignment, and multi-round proposal revisions.",
            "Add automated PDF proposal document ingestion and section-level novelty comparison."
        ])
    ]

    for title, left, top, bg_col, pts in future_cards:
        add_card(slide8, left, top, grid_w, grid_h, bg_col)
        tb = slide8.shapes.add_textbox(left + Inches(0.15), top + Inches(0.10), grid_w - Inches(0.3), grid_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Times New Roman"
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_ACCENT
        pt.space_after = Pt(4)
        for t in pts:
            p = tf.add_paragraph()
            p.text = f"• {t}"
            p.font.name = "Times New Roman"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(2)

    # Save presentation
    prs.save(OUTPUT_PPTX)
    print(f"Presentation successfully updated at: {OUTPUT_PPTX}")

if __name__ == "__main__":
    create_presentation()
